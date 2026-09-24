import os
import logging
import smtplib
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
import pandas as pd

from database import get_db
from config import REPORTS_DIR, BRANCH_METADATA, SENDER_EMAIL, SMTP_HOST, SMTP_PORT, SMTP_USER, SMTP_PASSWORD
from utils import get_year_label

logger = logging.getLogger("attendance.reports")

async def generate_attendance_excel(
    subject: str,
    start_date: str,
    end_date: str,
    section: str | None = None,
    branch: str | None = None,
    year: int | None = None,
) -> tuple[str, str]:
    """
    Query attendance for the specified subject between start_date and end_date, 
    join with master roster, and write a formatted Cumulative Excel register.
    Returns (filepath, filename).
    """
    clean_sec = section.strip().upper() if section else ""
    clean_branch = branch.strip().upper() if branch else ""

    db = await get_db()
    try:
        # Fetch Roster
        q_roster = """
            SELECT r.year, r.branch_code, r.branch_name, r.section, r.class_roll_no, r.primary_roll_no as roll_no, r.name,
                   (s.roll_no IS NOT NULL) as is_enrolled
            FROM college_roster r
            LEFT JOIN students s ON r.primary_roll_no = s.roll_no
            WHERE 1=1
        """
        params_roster = []
        if clean_sec:
            q_roster += " AND r.section = ?"
            params_roster.append(clean_sec)
        elif clean_branch:
            q_roster += " AND r.branch_code = ?"
            params_roster.append(clean_branch)
        if year:
            q_roster += " AND r.year = ?"
            params_roster.append(year)
        q_roster += " ORDER BY r.section, CAST(r.class_roll_no AS INTEGER), r.primary_roll_no"
        
        cur_r = await db.execute(q_roster, params_roster)
        roster_rows = await cur_r.fetchall()
        
        # Fetch Attendance
        q_att = """
            SELECT roll_no, date, status 
            FROM attendance 
            WHERE subject = ? AND date BETWEEN ? AND ?
        """
        cur_a = await db.execute(q_att, (subject, start_date, end_date))
        att_rows = await cur_a.fetchall()
    finally:
        await db.close()

    # Process Attendance into a dictionary: att_dict[roll_no][date] = "P" or "A"
    att_dict = {}
    dates_seen = set()
    for row in att_rows:
        r_no, dt, st = row["roll_no"], row["date"], row["status"]
        if r_no not in att_dict:
            att_dict[r_no] = {}
        # Convert "Present" -> "P", anything else/absent -> "A"
        att_dict[r_no][dt] = "P" if st.lower() == "present" else "A"
        dates_seen.add(dt)
        
    sorted_dates = sorted(list(dates_seen))
    
    # Build Final Data
    final_data = []
    for r in roster_rows:
        roll = r["roll_no"]
        row_dict = {
            "Year": get_year_label(r["year"]),
            "Branch": r["branch_code"],
            "Section": r["section"],
            "Class Roll": r["class_roll_no"],
            "AKTU Roll": roll,
            "Name": r["name"],
            "Biometric Reg": "Yes" if r["is_enrolled"] else "No",
        }
        
        total_p = 0
        total_classes = len(sorted_dates)
        for dt in sorted_dates:
            status = att_dict.get(roll, {}).get(dt, "A")
            row_dict[dt] = status
            if status == "P":
                total_p += 1
                
        row_dict["Total Present"] = total_p
        row_dict["Total Classes"] = total_classes
        row_dict["Percentage"] = f"{int((total_p / total_classes) * 100)}%" if total_classes > 0 else "0%"
        
        final_data.append(row_dict)

    df = pd.DataFrame(final_data)
    
    filename = f"Attendance_{subject.replace(' ', '_')}_{start_date}_to_{end_date}.xlsx"
    filepath = os.path.join(REPORTS_DIR, filename)
    
    # Save with formatting
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name="Attendance")
        worksheet = writer.sheets["Attendance"]
        
        # Color coding P=Green, A=Red
        from openpyxl.styles import PatternFill
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        
        for row in worksheet.iter_rows(min_row=2, min_col=8, max_col=7+len(sorted_dates)):
            for cell in row:
                if cell.value == "P":
                    cell.fill = green_fill
                elif cell.value == "A":
                    cell.fill = red_fill

    return filepath, filename



def send_email_with_attachment(
    to_email: str,
    subject_line: str,
    body: str,
    attachment_path: str,
) -> None:
    """
    Send an email with a single ``.xlsx`` attachment via SMTP (TLS).
    """
    msg = MIMEMultipart()
    msg["From"] = SENDER_EMAIL
    msg["To"] = to_email
    msg["Subject"] = subject_line

    msg.attach(MIMEText(body, "plain"))

    basename = os.path.basename(attachment_path)
    with open(attachment_path, "rb") as f:
        part = MIMEApplication(f.read(), Name=basename)
    part["Content-Disposition"] = f'attachment; filename="{basename}"'
    msg.attach(part)

    with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
        server.ehlo()
        server.starttls()
        server.login(SMTP_USER, SMTP_PASSWORD)
        server.send_message(msg)

    logger.info("Email sent to %s — subject: %s", to_email, subject_line)
